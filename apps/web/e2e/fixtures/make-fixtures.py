"""Builds the Excel files the bill-import browser flow uploads (run: python3 make-fixtures.py; needs openpyxl)."""
from datetime import date
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font

HERE = Path(__file__).parent

# A product list as a distributor keeps it in Excel.
wb = Workbook()
ws = wb.active
ws.title = "Items"
ws.append(["Item Name", "GST %", "HSN Code", "UOM", "Case Unit", "Case Size", "Company", "Sale Price", "Purchase Rate", "Barcode"])
ws.append(["Surf Excel 1 kg", 18, "3402", "PCS", "CTN", 12, "HUL", 128, 115, "8901030865278"])
ws.append(["Parle-G 100g", 5, "1905", "PCS", None, None, "Parle", 10, 8, "8901719101045"])
ws.append(["Tata Salt 1 kg", 0, "2501", "PCS", None, None, "Tata", 28, 22, None])
ws.append(["Harpic Toilet Cleaner 500ml", 18, "3402", "PCS", None, None, "Reckitt", 99, 85, None])
wb.save(HERE / "products.xlsx")

# A supplier's bill exported from their billing software: title rows, then the
# header, the bill number only on the first line, split CGST / SGST, a scheme
# (free) quantity, a real Excel date, and a total row at the end.
wb = Workbook()
ws = wb.active
ws.title = "Sale Bill"
ws.append(["SHREE GANESH AGENCIES"])
ws["A1"].font = Font(bold=True, size=14)
ws.append(["12 Siyaganj, Indore · GSTIN: 23AAACM1234M1ZR"])
ws.append([])
ws.append(["Invoice No.", "Inv Date", "Party GSTIN", "Particulars", "Item Code", "Barcode", "HSN", "Qty", "Unit", "Sch Qty", "PTR", "Disc %", "CGST %", "SGST %", "Batch", "Exp.", "M.R.P.", "Amount"])
rows = [
    ["SG/1021", date(2026, 9, 20), "23AAACM1234M1ZR", "SURF EXCEL EASY WASH 1KG", "HU-1102", "8901030865278", "3402", 2, "CS", None, 1380, 2.5, 9, 9, None, None, 140, 2691],
    [None, None, None, "PARLE-G BISCUIT 100 GM", "PA-100", "8901719101045", "1905", 240, "PCS", 12, 8, None, 2.5, 2.5, None, None, 10, 1920],
    [None, None, None, "TATA SALT 1KG", "TS-1", None, "2501", 50, "PCS", None, 22, None, 0, 0, None, None, 28, 1100],
    [None, None, None, "HARPIC 500ML", "HP-500", None, "3402", 12, "PCS", None, 85, None, 9, 9, None, None, 99, 1020],
    [None, None, None, "COLGATE STRONG TEETH 200G", "CG-200", "8901314010322", "3306", 24, "PCS", None, 92, None, 9, 9, None, None, 110, 2208],
]
for r in rows:
    ws.append(r)
ws["B5"].number_format = "dd/mm/yyyy"
ws.append([None, None, None, "Total", None, None, None, None, None, None, None, None, None, None, None, None, None, 8939])
wb.save(HERE / "supplier-bill.xlsx")
print("written")
